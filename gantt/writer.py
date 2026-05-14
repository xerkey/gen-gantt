from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .layout import (
    COL_ASSIGNEE,
    COL_DATE_START,
    COL_END,
    COL_NAME,
    COL_ROLE,
    COL_START,
    ROW_DAY,
    ROW_MONTH,
    ROW_TASK_START,
    ROW_WEEK,
    Layout,
    week_label,
)


HEADER_FILL = PatternFill("solid", fgColor="FFE6E6E6")
MONTH_FILLS = (
    PatternFill("solid", fgColor="FFD9D9D9"),
    PatternFill("solid", fgColor="FFB7B7B7"),
)
WEEK_FILLS = (
    PatternFill("solid", fgColor="FFEDEDED"),
    PatternFill("solid", fgColor="FFCFCFCF"),
)
SATURDAY_FILL = PatternFill("solid", fgColor="FFCCE5FF")
SUNDAY_FILL = PatternFill("solid", fgColor="FFFFCCCC")
BAR_FILLS = {
    0: PatternFill("solid", fgColor="FF1F4E78"),
    1: PatternFill("solid", fgColor="FF4472C4"),
    2: PatternFill("solid", fgColor="FFB4C7E7"),
}
LEVEL_FONTS = {
    0: Font(bold=True),
    1: Font(),
    2: Font(color="FF808080"),
}
HEADER_FONT = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")


def _write_headers(ws: Worksheet, layout: Layout) -> None:
    label_columns = [
        (COL_ROLE, "役割", LEFT),
        (COL_ASSIGNEE, "担当者", LEFT),
        (COL_NAME, "タスク名", LEFT),
        (COL_START, "開始日", CENTER),
        (COL_END, "終了日", CENTER),
    ]
    for col, label, align in label_columns:
        cell = ws.cell(row=ROW_DAY, column=col, value=label)
        cell.font = HEADER_FONT
        cell.alignment = align

    for r in (ROW_MONTH, ROW_WEEK, ROW_DAY):
        for col, _, _ in label_columns:
            ws.cell(row=r, column=col).fill = HEADER_FILL

    prev_month: tuple[int, int] | None = None
    month_idx = -1
    week_idx = -1
    for dc in layout.date_columns:
        d = dc.date
        month_key = (d.year, d.month)
        if month_key != prev_month:
            month_idx += 1
            ws.cell(row=ROW_MONTH, column=dc.col, value=f"{d.year}-{d.month:02d}")
            prev_month = month_key

        if dc.is_week_start:
            week_idx += 1
            ws.cell(
                row=ROW_WEEK,
                column=dc.col,
                value=week_label(d, layout.project.week_start),
            )

        ws.cell(row=ROW_DAY, column=dc.col, value=d.day)

        month_fill = MONTH_FILLS[month_idx % len(MONTH_FILLS)]
        week_fill = WEEK_FILLS[week_idx % len(WEEK_FILLS)]
        weekday = d.weekday()
        if weekday == 5:
            day_fill = SATURDAY_FILL
        elif weekday == 6:
            day_fill = SUNDAY_FILL
        else:
            day_fill = HEADER_FILL

        ws.cell(row=ROW_MONTH, column=dc.col).fill = month_fill
        ws.cell(row=ROW_WEEK, column=dc.col).fill = week_fill
        ws.cell(row=ROW_DAY, column=dc.col).fill = day_fill

        for r in (ROW_MONTH, ROW_WEEK, ROW_DAY):
            cell = ws.cell(row=r, column=dc.col)
            cell.alignment = CENTER
            cell.font = HEADER_FONT


def _write_task_rows(ws: Worksheet, layout: Layout) -> None:
    for tr in layout.task_rows:
        task = tr.task
        font = LEVEL_FONTS[task.level]

        role_cell = ws.cell(row=tr.row, column=COL_ROLE, value=task.role or None)
        role_cell.font = font
        role_cell.alignment = LEFT

        assignee_cell = ws.cell(row=tr.row, column=COL_ASSIGNEE, value=task.assignee or None)
        assignee_cell.font = font
        assignee_cell.alignment = LEFT

        indent = "  " * task.level
        name_cell = ws.cell(row=tr.row, column=COL_NAME, value=f"{indent}{task.name}")
        name_cell.font = font
        name_cell.alignment = LEFT

        if task.start is not None:
            start_cell = ws.cell(row=tr.row, column=COL_START, value=task.start)
            start_cell.number_format = "yyyy-mm-dd"
            start_cell.alignment = CENTER
            start_cell.font = font
        if task.end is not None:
            end_cell = ws.cell(row=tr.row, column=COL_END, value=task.end)
            end_cell.number_format = "yyyy-mm-dd"
            end_cell.alignment = CENTER
            end_cell.font = font


def _draw_bars(ws: Worksheet, layout: Layout) -> None:
    for tr in layout.task_rows:
        task = tr.task
        if task.start is None or task.end is None:
            continue
        fill = BAR_FILLS[task.level]
        start_col = layout.col_for_date(task.start)
        end_col = layout.col_for_date(task.end)
        for c in range(start_col, end_col + 1):
            ws.cell(row=tr.row, column=c).fill = fill


def _apply_dimensions(ws: Worksheet, layout: Layout) -> None:
    ws.column_dimensions[get_column_letter(COL_ROLE)].width = 16
    ws.column_dimensions[get_column_letter(COL_ASSIGNEE)].width = 16
    ws.column_dimensions[get_column_letter(COL_NAME)].width = 30
    ws.column_dimensions[get_column_letter(COL_START)].width = 12
    ws.column_dimensions[get_column_letter(COL_END)].width = 12

    for dc in layout.date_columns:
        letter = get_column_letter(dc.col)
        cd = ws.column_dimensions[letter]
        cd.width = 3
        if dc.outline_level > 0:
            cd.outline_level = dc.outline_level
            cd.hidden = False


def _apply_row_outline(ws: Worksheet, layout: Layout) -> None:
    for tr in layout.task_rows:
        if tr.outline_level > 0:
            rd = ws.row_dimensions[tr.row]
            rd.outline_level = tr.outline_level
            rd.hidden = False


def _apply_outline_properties(ws: Worksheet) -> None:
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_properties.outlinePr.summaryRight = False


def _freeze_panes(ws: Worksheet) -> None:
    ws.freeze_panes = ws.cell(row=ROW_TASK_START, column=COL_DATE_START).coordinate


def write_workbook(layout: Layout, output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Gantt"

    _apply_outline_properties(ws)
    _write_headers(ws, layout)
    _write_task_rows(ws, layout)
    _draw_bars(ws, layout)
    _apply_dimensions(ws, layout)
    _apply_row_outline(ws, layout)
    _freeze_panes(ws)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
